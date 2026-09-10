import os, io, math, shutil, subprocess
from datetime import datetime
import chess, chess.pgn, chess.engine
import gradio as gr
import matplotlib
matplotlib.use("Agg")
from matplotlib.figure import Figure
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

if not shutil.which("stockfish"):
    subprocess.run(["apt-get","update","-qq"], check=False)
    subprocess.run(["apt-get","install","-y","-qq","stockfish"], check=False)
STOCKFISH_PATH = shutil.which("stockfish") or "/usr/games/stockfish"

VALORES_PECAS = {chess.PAWN:1, chess.KNIGHT:3, chess.BISHOP:3, chess.ROOK:5, chess.QUEEN:9, chess.KING:0}

def detectar_plataforma(h):
    s = (h.get("Site","") or "").lower(); e = (h.get("Event","") or "").lower()
    if "chess.com" in s or "chess.com" in e: return "Chess.com"
    if "lichess" in s or "lichess" in e: return "Lichess"
    return "Desconhecida"

def calc_material(board):
    t = 0
    for casa in chess.SQUARES:
        p = board.piece_at(casa)
        if p:
            v = VALORES_PECAS.get(p.piece_type, 0)
            t += v if p.color == chess.WHITE else -v
    return t

def validar_pgn(texto):
    if not texto or not texto.strip(): raise ValueError("Cole um PGN.")
    game = chess.pgn.read_game(io.StringIO(texto))
    if game is None: raise ValueError("PGN invalido.")
    lances = list(game.mainline_moves())
    if not lances: raise ValueError("PGN sem lances.")
    board = game.board()
    for i, m in enumerate(lances, 1):
        if not board.is_legal(m): raise ValueError(f"Lance ilegal {i}")
        board.push(m)
    return game

def cp_prob(cp):
    cp = max(-1500, min(1500, cp)); return 1.0/(1.0+math.pow(10,-cp/400.0))

def prec_lance(cp_a, cp_d, cor):
    if cor == "brancas": a, d = cp_a, cp_d
    else: a, d = -cp_a, -cp_d
    perda = max(0, cp_prob(a) - cp_prob(d))
    return max(0, min(100, 100*(1-perda*2.2)))

def classificar(p, san, melhor):
    if "#" in san: return "Melhor", "Xeque-mate"
    if san == melhor: return "Melhor", "Melhor lance"
    if p >= 95: return "Excelente", "Alto nivel"
    if p >= 85: return "Bom", "Solido"
    if p >= 70: return "Imprecisao", "Poderia ser melhor"
    if p >= 50: return "Erro", "Erro tatico"
    return "Erro grave", "Erro grave"

def analisar_stockfish(game, prof=10):
    if not STOCKFISH_PATH: return None
    try:
        eng = chess.engine.SimpleEngine.popen_uci(STOCKFISH_PATH)
        try: eng.configure({"Threads":2,"Hash":128})
        except: pass
        board = game.board(); res = []
        for move in list(game.mainline_moves()):
            info = eng.analyse(board, chess.engine.Limit(depth=prof))
            sc = info["score"].pov(chess.WHITE)
            cp = sc.score(mate_score=100000) or 0
            pv = info.get("pv") or []
            melhor = board.san(pv[0]) if pv else ""
            san = board.san(move)
            board.push(move)
            res.append({"cp":cp,"melhor_san":melhor,"san":san})
        eng.quit(); return res
    except Exception as e:
        print(f"Stockfish: {e}"); return None

def analisar(game):
    headers = dict(game.headers); lances = list(game.mainline_moves()); board = game.board()
    cab = {"brancas": headers.get("White","Brancas"),"negras": headers.get("Black","Pretas"),
           "rating_brancas": headers.get("WhiteElo","-"),"rating_negras": headers.get("BlackElo","-"),
           "resultado": headers.get("Result","*"),"evento": headers.get("Event","-"),
           "data": headers.get("Date","-"),"plataforma": detectar_plataforma(headers),
           "total_lances": len(lances)}
    eng_res = analisar_stockfish(game) if STOCKFISH_PATH else None
    dados = []; stats = {"brancas":{"classes":{},"precs":[]},"negras":{"classes":{},"precs":[]}}
    curva = [0]; marcadores = []
    for i, move in enumerate(lances):
        cor = "brancas" if i%2==0 else "negras"
        san = board.san(move)
        board.push(move)
        prec = 100.0; melhor = san; ev = 0
        if eng_res and i < len(eng_res):
            info = eng_res[i]; ev = info["cp"]; melhor = info["melhor_san"]
            cp_a = eng_res[i-1]["cp"] if i>0 else 20
            prec = prec_lance(cp_a, ev, cor)
        else:
            prec = 75
        cls, motivo = classificar(prec, san, melhor)
        cp_vis = max(-1000, min(1000, ev))
        curva.append(cp_vis/100)
        if cls in ("Erro","Erro grave","Imprecisao"):
            marcadores.append({"ply":i+1,"cp":cp_vis/100,"classe":cls})
        dados.append({"n":i+1,"cor":cor,"san":san,"classe":cls,"motivo":motivo,
                      "precisao":prec,"eval":ev,"melhor":melhor})
        stats[cor]["classes"][cls] = stats[cor]["classes"].get(cls,0)+1
        stats[cor]["precs"].append(prec)
    pb = sum(stats["brancas"]["precs"])/max(1,len(stats["brancas"]["precs"]))
    pn = sum(stats["negras"]["precs"])/max(1,len(stats["negras"]["precs"]))
    def nota(p):
        return "A+" if p>=95 else "A" if p>=90 else "B+" if p>=85 else "B" if p>=80 else "C+" if p>=75 else "C" if p>=70 else "D" if p>=60 else "F"
    return {"cabecalho":cab,"lances":dados,"estatisticas":stats,
            "precisao_brancas":pb,"precisao_negras":pn,
            "nota_brancas":nota(pb),"nota_negras":nota(pn),
            "curva":curva,"marcadores":marcadores}

def fig_evolucao(d):
    c = d["curva"]; xs = list(range(len(c)))
    fig = Figure(figsize=(9,4),dpi=110,facecolor="#0d1117")
    ax = fig.add_subplot(111,facecolor="#0d1117")
    ax.fill_between(xs,[max(0,y) for y in c],0,color="#f0f6fc",alpha=0.92,zorder=1)
    ax.fill_between(xs,[min(0,y) for y in c],0,color="#0d1117",alpha=0.95,zorder=1)
    ax.plot(xs,c,color="#58a6ff",linewidth=1.4,zorder=3)
    ax.axhline(0,color="#30363d",linewidth=1,zorder=2)
    for cl,cor in {"Imprecisao":"#d29922","Erro":"#f0883e","Erro grave":"#f85149"}.items():
        pts = [m for m in d["marcadores"] if m["classe"]==cl]
        if pts:
            ax.scatter([p["ply"] for p in pts],[p["cp"] for p in pts],color=cor,s=65,zorder=5,
                       edgecolors="#0d1117",linewidths=1.2,label=cl)
    ax.set_xlim(0,max(1,len(c)-1)); ax.set_ylim(-8,8)
    ax.tick_params(colors="#8b949e",labelsize=8)
    ax.grid(True,color="#21262d",linestyle=":",alpha=0.5)
    for sp in ax.spines.values(): sp.set_color("#30363d")
    leg = ax.legend(loc="upper right",fontsize=7,facecolor="#161b22",edgecolor="#30363d")
    if leg:
        for t in leg.get_texts(): t.set_color("#c9d1d9")
    fig.tight_layout(); return fig

def html_resultado(d):
    c = d["cabecalho"]; pb,pn = d["precisao_brancas"],d["precisao_negras"]
    nb,nn = d["nota_brancas"],d["nota_negras"]
    def cor(n): return {"A+":"#3fb950","A":"#3fb950","B+":"#58a6ff","B":"#58a6ff","C+":"#d29922","C":"#d29922","D":"#f0883e","F":"#f85149"}.get(n,"#c9d1d9")
    stats_html = ""
    for cat in ["Melhor","Excelente","Bom","Imprecisao","Erro","Erro grave"]:
        cb = d["estatisticas"]["brancas"]["classes"].get(cat,0)
        cn = d["estatisticas"]["negras"]["classes"].get(cat,0)
        stats_html += f'<tr><td style="padding:8px;text-align:center;color:#58a6ff;">{cb}</td><td style="padding:8px;text-align:center;font-weight:600;">{cat}</td><td style="padding:8px;text-align:center;color:#f0883e;">{cn}</td></tr>'
    lances_html = ""
    for l in d["lances"][:60]:
        bg = "#1c2128" if l["cor"]=="brancas" else "#161b22"
        lances_html += f'<tr style="background:{bg};"><td style="padding:6px 8px;color:#8b949e;font-size:12px;">{l["n"]}</td><td style="padding:6px 8px;font-weight:600;">{l["san"]}</td><td style="padding:6px 8px;font-size:12px;">{l["classe"]}</td><td style="padding:6px 8px;text-align:right;color:#58a6ff;">{l["precisao"]:.0f}%</td></tr>'
    return f'<div style="background:#0d1117;color:#c9d1d9;padding:16px;border-radius:14px;font-family:sans-serif;"><h2 style="color:#58a6ff;text-align:center;">Resultado</h2><div style="display:grid;grid-template-columns:1fr 1fr;gap:10px;margin:14px 0;"><div style="background:#161b22;border:1px solid #30363d;border-radius:12px;padding:14px;text-align:center;"><div style="color:#8b949e;font-size:11px;">BRANCAS</div><div style="font-weight:600;">{c["brancas"]}</div><div style="font-size:34px;font-weight:700;color:{cor(nb)};">{pb:.1f}%</div><div style="display:inline-block;padding:3px 10px;border-radius:14px;background:{cor(nb)};color:#fff;font-weight:600;">Nota {nb}</div></div><div style="background:#161b22;border:1px solid #30363d;border-radius:12px;padding:14px;text-align:center;"><div style="color:#8b949e;font-size:11px;">NEGRAS</div><div style="font-weight:600;">{c["negras"]}</div><div style="font-size:34px;font-weight:700;color:{cor(nn)};">{pn:.1f}%</div><div style="display:inline-block;padding:3px 10px;border-radius:14px;background:{cor(nn)};color:#fff;font-weight:600;">Nota {nn}</div></div></div><h3 style="color:#58a6ff;">Estatisticas</h3><table style="width:100%;border-collapse:collapse;font-size:12px;"><tr style="background:#21262d;"><th style="padding:8px;">{c["brancas"][:10]}</th><th style="padding:8px;">Categoria</th><th style="padding:8px;">{c["negras"][:10]}</th></tr>{stats_html}</table><h3 style="color:#58a6ff;">Historico</h3><div style="max-height:280px;overflow-y:auto;"><table style="width:100%;border-collapse:collapse;font-size:12px;"><tr style="background:#21262d;"><th style="padding:6px 8px;text-align:left;">#</th><th style="padding:6px 8px;text-align:left;">Lance</th><th style="padding:6px 8px;text-align:left;">Classe</th><th style="padding:6px 8px;text-align:right;">Prec.</th></tr>{lances_html}</table></div></div>'

def gerar_word(d, caminho="/tmp/relatorio.docx"):
    doc = Document()
    t = doc.add_heading("E4 Analyzer - Relatorio", 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = d["cabecalho"]
    doc.add_paragraph(f"Brancas: {c['brancas']} ({c['rating_brancas']}) - {d['precisao_brancas']:.1f}% Nota {d['nota_brancas']}")
    doc.add_paragraph(f"Negras: {c['negras']} ({c['rating_negras']}) - {d['precisao_negras']:.1f}% Nota {d['nota_negras']}")
    doc.add_paragraph(f"Resultado: {c['resultado']}")
    doc.add_heading("Historico", 1)
    tab = doc.add_table(rows=1, cols=4); tab.style = "Light Grid Accent 1"
    h = tab.rows[0].cells
    h[0].text="#"; h[1].text="Lance"; h[2].text="Classe"; h[3].text="Precisao"
    for l in d["lances"]:
        r = tab.add_row().cells
        r[0].text=str(l["n"]); r[1].text=l["san"]; r[2].text=l["classe"]; r[3].text=f"{l['precisao']:.0f}%"
    doc.save(caminho); return caminho

def processar(pgn_text):
    try:
        game = validar_pgn(pgn_text)
        d = analisar(game)
        return html_resultado(d), fig_evolucao(d), gerar_word(d), "Analise concluida!"
    except Exception as e:
        return f'<div style="color:#f85149;padding:20px;">Erro: {e}</div>', None, None, f"Erro: {e}"

with gr.Blocks(title="E4 Analyzer") as app:
    gr.Markdown("# E4 Analyzer\n**Analise profissional de partidas**")
    pgn_in = gr.Textbox(label="Cole o PGN", lines=10)
    with gr.Row():
        btn_analisar = gr.Button("ANALISAR", variant="primary", size="lg")
        btn_limpar = gr.Button("LIMPAR", variant="secondary", size="lg")
    msg = gr.Markdown("")
    html_out = gr.HTML()
    grafico_out = gr.Plot(label="Evolucao")
    arquivo_word = gr.File(label="Baixar Word")
    btn_analisar.click(fn=processar, inputs=[pgn_in], outputs=[html_out, grafico_out, arquivo_word, msg])
    btn_limpar.click(fn=lambda: ("", None, None, None, ""), outputs=[pgn_in, html_out, grafico_out, arquivo_word, msg])

if __name__ == "__main__":
    app.launch(server_name="0.0.0.0", server_port=int(os.environ.get("PORT", 7860)))
